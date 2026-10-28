"""Generazione del file Excel per il carico massivo nel gestionale.

Il primo foglio contiene esattamente i campi richiesti dalla specifica
(Magazzino, Codice Articolo, Variante, Quantita) e nient'altro: e' il file da
dare in pasto all'import. Il secondo foglio e' la tracciabilita', per capire a
mesi di distanza da quale pagina e da quale riga di bolla arriva un carico. Il
terzo e' il registro delle modifiche: chi ha caricato la bolla, chi ha corretto
cosa e chi ha esportato.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXCEL_HEADERS, EXPORT_DIR, MAGAZZINI

_INTESTAZIONE = Font(bold=True, color="FFFFFF")
_SFONDO = PatternFill("solid", fgColor="2F5233")

TRACCIABILITA = [
    "Pagina",
    "Seq",
    "Codice Articolo",
    "Descrizione",
    "Colore",
    "Taglia bolla",
    "Variante",
    "Quantita",
    "Magazzino",
    "Origine",
    "Note",
    # l'operatore ha controllato gli avvisi sulla scansione (chi, nel foglio Modifiche)
    "Verificata",
    # valori stampati sostituiti da una correzione a mano; vuoti se nessuna
    "Codice stampato",
    "Taglia stampata",
    "Qta stampata",
    # chi ha caricato la bolla e chi ha toccato la riga (vedi il foglio Modifiche)
    "Caricata da",
    "Modificata da",
]

MODIFICHE = [
    "Quando",
    "Operatore",
    "Azione",
    "Pagina",
    "Seq",
    "Codice Articolo",
    "Taglia",
    "Campo",
    "Prima",
    "Dopo",
]


def righe_da_esportare(righe: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Solo le righe attive, con quantita' e codice valorizzati."""
    return [
        r
        for r in righe
        if not r.get("eliminata") and int(r.get("qta") or 0) > 0 and r.get("codice_articolo")
    ]


def problemi_bloccanti(righe: list[dict[str, Any]]) -> list[str]:
    """Motivi per cui l'export non deve partire."""
    attive = righe_da_esportare(righe)
    problemi: list[str] = []
    if not attive:
        problemi.append("nessuna riga da esportare")
    senza_codice = [
        r for r in righe if not r.get("eliminata") and not r.get("codice_articolo")
    ]
    if senza_codice:
        problemi.append(
            f"{len(senza_codice)} righe attive senza codice articolo: "
            "completale o eliminale"
        )
    senza_variante = [r for r in attive if not r.get("variante")]
    if senza_variante:
        problemi.append(f"{len(senza_variante)} righe attive senza variante")
    senza_magazzino = [r for r in attive if not r.get("magazzino")]
    if senza_magazzino:
        problemi.append(f"{len(senza_magazzino)} righe attive senza magazzino")
    # un magazzino tolto da BOLLE_MAGAZZINI dopo la revisione non va nel gestionale
    sconosciuti = sorted({r["magazzino"] for r in attive if r.get("magazzino")} - set(MAGAZZINI))
    if sconosciuti:
        problemi.append(
            f"magazzino sconosciuto ({', '.join(sconosciuti)}): scegli fra {', '.join(MAGAZZINI)}"
        )
    # la variante di una taglia fuori da config/varianti.csv e' solo la taglia
    # ricopiata: finche' l'operatore non la conferma non va nel gestionale
    da_validare = [r for r in attive if r.get("variante") and not r.get("variante_confermata")]
    if da_validare:
        taglie = sorted({r.get("taglia") or r["variante"] for r in da_validare})
        problemi.append(
            f"{len(da_validare)} righe con la taglia da validare ({', '.join(taglie)}): "
            "conferma o correggi la variante"
        )
    return problemi


def aggrega(righe: list[dict[str, Any]]) -> list[tuple[str, str, str, int]]:
    """Somma le quantita' di (magazzino, codice, variante) identici.

    Capita spesso: lo stesso articolo compare su piu' righe della stessa bolla
    perche' arriva da ordini diversi (vedi le due righe TABELA 23 JSY taglia L
    a pagina 9 degli esempi). Il gestionale preferisce una riga sola.
    """
    somme: dict[tuple[str, str, str], int] = {}
    ordine: list[tuple[str, str, str]] = []
    for r in righe_da_esportare(righe):
        chiave = (r["magazzino"], r["codice_articolo"], r["variante"])
        if chiave not in somme:
            somme[chiave] = 0
            ordine.append(chiave)
        somme[chiave] += int(r["qta"])
    return [(m, c, v, somme[(m, c, v)]) for m, c, v in ordine]


def _intesta(ws, intestazioni: list[str]) -> None:
    ws.append(intestazioni)
    for cella in ws[1]:
        cella.font = _INTESTAZIONE
        cella.fill = _SFONDO
        cella.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autoampiezza(ws) -> None:
    for colonna in ws.columns:
        larghezza = max((len(str(c.value or "")) for c in colonna), default=8)
        ws.column_dimensions[get_column_letter(colonna[0].column)].width = min(
            max(larghezza + 2, 10), 45
        )


def _solo_testo(ws) -> None:
    """Niente formule: openpyxl scrive come formula ogni testo che inizia con '='.

    I testi arrivano dalla bolla (letta da un modello) o da chi usa l'app: un
    "=HYPERLINK(...)" o "=cmd|..." diventerebbe una formula eseguita da chi apre
    il file. Il valore resta identico, cambia solo il tipo della cella.
    """
    for riga in ws.iter_rows():
        for cella in riga:
            if cella.data_type == "f":
                cella.data_type = "s"


def esporta(doc: dict[str, Any], dest: Path | None = None) -> Path:
    righe = doc.get("righe", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Carico"
    _intesta(ws, EXCEL_HEADERS)
    for magazzino, codice, variante, qta in aggrega(righe):
        ws.append([magazzino, codice, variante, qta])
    _solo_testo(ws)
    _autoampiezza(ws)

    storico = doc.get("storico") or []
    autori: dict[str, list[str]] = {}
    for voce in storico:
        if voce.get("riga_id"):
            nomi = autori.setdefault(voce["riga_id"], [])
            if voce["operatore"] not in nomi:
                nomi.append(voce["operatore"])

    tracc = wb.create_sheet("Tracciabilita")
    _intesta(tracc, TRACCIABILITA)
    for r in righe:
        tracc.append(
            [
                r.get("pagina"),
                r.get("seq"),
                r.get("codice_articolo"),
                r.get("descrizione"),
                r.get("colore"),
                r.get("taglia"),
                r.get("variante"),
                r.get("qta"),
                r.get("magazzino"),
                "ELIMINATA" if r.get("eliminata") else r.get("origine"),
                " | ".join(
                    filter(
                        None,
                        [r.get("annotazioni_manoscritte", ""), *(r.get("avvisi") or [])],
                    )
                ),
                "si'" if r.get("verificata") else "",
                r.get("codice_stampato", ""),
                r.get("taglia_stampata", ""),
                r["qta_stampata"] if r.get("qta_stampata", -1) >= 0 else "",
                doc.get("operatore", ""),
                ", ".join(autori.get(r.get("id", ""), [])),
            ]
        )
    _solo_testo(tracc)
    _autoampiezza(tracc)

    modifiche = wb.create_sheet("Modifiche")
    _intesta(modifiche, MODIFICHE)
    for voce in storico:
        modifiche.append([voce.get(k, "") for k in (
            "quando", "operatore", "azione", "pagina", "seq", "codice_articolo",
            "taglia", "campo", "prima", "dopo",
        )])
    _solo_testo(modifiche)
    _autoampiezza(modifiche)

    if dest is None:
        stampo = datetime.now().strftime("%Y%m%d-%H%M%S")
        dest = EXPORT_DIR / f"carico_{doc.get('id', 'doc')}_{stampo}.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest
